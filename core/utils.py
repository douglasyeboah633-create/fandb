"""Helper utilities: exports (CSV / Excel / PDF), IP address and formatting."""

import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.conf import settings
from django.http import HttpResponse
from django.utils import timezone


# ---------------------------------------------------------------------------
# Generic helpers
# ---------------------------------------------------------------------------

def get_client_ip(request):
    """Return the best-effort client IP address for audit logging."""
    forwarded = request.META.get("HTTP_X_FORWARDED_FOR")
    if forwarded:
        return forwarded.split(",")[0].strip()
    return request.META.get("REMOTE_ADDR")


def money(value) -> str:
    """Format a value as a currency amount (thousands separated)."""
    if value in (None, ""):
        value = 0
    try:
        amount = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return str(value)
    return f"{amount:,.2f}"


def currency(value) -> str:
    """Format a value with the configured currency symbol."""
    return f"{settings.CURRENCY_SYMBOL} {money(value)}"


def parse_date(value, default=None):
    """Parse a YYYY-MM-DD string into a date, returning default on failure."""
    if not value:
        return default
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return default


def clean_filename(name: str) -> str:
    """Make a string safe to use as a download filename."""
    safe = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in str(name))
    return safe.strip("_") or "export"


def _timestamp() -> str:
    return timezone.localtime().strftime("%Y%m%d_%H%M")


# ---------------------------------------------------------------------------
# CSV export (opens directly in Excel)
# ---------------------------------------------------------------------------

def csv_response(filename: str, headers, rows) -> HttpResponse:
    """Build an HTTP response containing a CSV file."""
    base = clean_filename(filename)
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = (
        f'attachment; filename="{base}_{_timestamp()}.csv"'
    )
    # UTF-8 BOM so Excel opens the file with the correct encoding.
    response.write("\ufeff")
    writer = csv.writer(response)
    writer.writerow(list(headers))
    for row in rows:
        writer.writerow(["" if value is None else value for value in row])
    return response


# ---------------------------------------------------------------------------
# Excel (.xlsx) export
# ---------------------------------------------------------------------------

def xlsx_response(filename: str, title: str, headers, rows) -> HttpResponse:
    """Build an HTTP response containing a formatted Excel workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (clean_filename(title) or "Report")[:31]

    header_fill = PatternFill("solid", fgColor="1F5C3A")
    header_font = Font(bold=True, color="FFFFFF")

    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append(["" if value is None else value for value in row])

    for index, header in enumerate(headers, start=1):
        length = max(len(str(header)) + 4, 12)
        sheet.column_dimensions[get_column_letter(index)].width = min(length, 42)

    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)

    response = HttpResponse(
        stream.read(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    base = clean_filename(filename)
    response["Content-Disposition"] = (
        f'attachment; filename="{base}_{_timestamp()}.xlsx"'
    )
    return response


# ---------------------------------------------------------------------------
# PDF export (landscape report tables)
# ---------------------------------------------------------------------------

def _pdf_styles():
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

    base = getSampleStyleSheet()
    return {
        "colors": colors,
        "business": ParagraphStyle(
            "BusinessName", parent=base["Title"], fontSize=15, leading=19,
            textColor=colors.HexColor("#1F5C3A"), spaceAfter=2,
        ),
        "title": ParagraphStyle(
            "ReportTitle", parent=base["Heading2"], fontSize=12, leading=16,
            textColor=colors.HexColor("#1A1A1A"), spaceAfter=2,
        ),
        "meta": ParagraphStyle(
            "Meta", parent=base["Normal"], fontSize=8, leading=11,
            textColor=colors.HexColor("#666666"),
        ),
        "summary": ParagraphStyle(
            "Summary", parent=base["Normal"], fontSize=9, leading=12,
            textColor=colors.HexColor("#1A1A1A"),
        ),
        "cell": ParagraphStyle("Cell", parent=base["Normal"], fontSize=7.5, leading=9.5),
        "head": ParagraphStyle(
            "HeadCell", parent=base["Normal"], fontSize=7.5, leading=9.5,
            textColor=colors.white, fontName="Helvetica-Bold",
        ),
    }


def _pdf_header(styles, title: str, subtitle: str = ""):
    from reportlab.platypus import Paragraph, Spacer

    contact = (
        f"{settings.BUSINESS_PHONE} &nbsp;|&nbsp; {settings.BUSINESS_EMAIL} "
        f"&nbsp;|&nbsp; {settings.BUSINESS_ADDRESS}"
    )
    generated = timezone.localtime().strftime("%d %b %Y %H:%M")
    return [
        Paragraph(settings.BUSINESS_NAME, styles["business"]),
        Paragraph(title, styles["title"]),
        Paragraph(contact, styles["meta"]),
        Paragraph(f"Generated: {generated}. {subtitle}", styles["meta"]),
        Spacer(1, 8),
    ]


def _pdf_table(styles, headers, rows, numeric_columns=None):
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Table, TableStyle

    numeric_columns = numeric_columns or set()
    data = [[Paragraph(str(h), styles["head"]) for h in headers]]
    for row in rows:
        data.append(
            [Paragraph("" if v is None else str(v), styles["cell"]) for v in row]
        )

    table = Table(data, repeatRows=1, hAlign="LEFT")
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F5C3A")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7C7BC")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F7F4")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for column in numeric_columns:
        style.append(("ALIGN", (column, 0), (column, -1), "RIGHT"))
    table.setStyle(TableStyle(style))
    return table


def pdf_response(
    filename: str,
    title: str,
    headers,
    rows,
    subtitle: str = "",
    summary: str = "",
    numeric_columns=None,
) -> HttpResponse:
    """Build a printable landscape PDF report."""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = _pdf_styles()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm,
        topMargin=14 * mm, bottomMargin=14 * mm,
        title=title, author=settings.BUSINESS_NAME,
    )

    flowables = _pdf_header(styles, title, subtitle)
    if summary:
        flowables.append(Paragraph(summary, styles["summary"]))
        flowables.append(Spacer(1, 6))
    flowables.append(_pdf_table(styles, headers, rows, numeric_columns))

    doc.build(flowables)
    buffer.seek(0)

    response = HttpResponse(buffer.read(), content_type="application/pdf")
    base = clean_filename(filename)
    response["Content-Disposition"] = (
        f'attachment; filename="{base}_{_timestamp()}.pdf"'
    )
    return response


def receipt_pdf_response(payment, business: dict) -> HttpResponse:
    """Build a single page payment receipt PDF."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = _pdf_styles()
    sale = payment.sale
    customer = sale.customer
    land = sale.land

    total_paid = sale.total_paid
    remaining = sale.outstanding_balance
    previous_balance = total_paid - payment.amount
    symbol = business.get("currency", settings.CURRENCY_SYMBOL)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=16 * mm, bottomMargin=16 * mm,
        title=f"Receipt {payment.receipt_number}",
        author=business.get("name", settings.BUSINESS_NAME),
    )

    recorded_by = "-"
    if payment.recorded_by:
        recorded_by = payment.recorded_by.get_full_name() or payment.recorded_by.username

    story = [
        Paragraph(business.get("name", settings.BUSINESS_NAME), styles["business"]),
        Paragraph("PAYMENT RECEIPT", styles["title"]),
        Paragraph(
            f"{business.get('address', '')} &nbsp;|&nbsp; {business.get('phone', '')}"
            f" &nbsp;|&nbsp; {business.get('email', '')}",
            styles["meta"],
        ),
        Spacer(1, 10),
    ]

    info_rows = [
        ["Receipt Number", payment.receipt_number,
         "Payment Date", payment.payment_date.strftime("%d %b %Y")],
        ["Customer", customer.full_name, "Customer Phone", customer.phone],
        ["Plot / Land No.", f"{land.plot_number} ({land.land_id})",
         "Location", land.location],
        ["Payment Method", payment.method_display, "Recorded By", recorded_by],
        ["Transaction ID", sale.transaction_id, "Sale Date",
         sale.sale_date.strftime("%d %b %Y")],
    ]

    body = styles["cell"]
    info_table = Table(
        [
            [
                Paragraph(f"<b>{a}</b>", body), Paragraph(str(b), body),
                Paragraph(f"<b>{c}</b>", body), Paragraph(str(d), body),
            ]
            for a, b, c, d in info_rows
        ],
        colWidths=[32 * mm, 55 * mm, 32 * mm, 55 * mm],
    )
    info_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7C7BC")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F2F7F4")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F2F7F4")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story += [info_table, Spacer(1, 12)]

    money_rows = [
        ["Selling Price", f"{symbol} {money(sale.selling_price)}"],
        ["Other Payments Recorded (including later payments)", f"{symbol} {money(previous_balance)}"],
        ["Amount Paid", f"{symbol} {money(payment.amount)}"],
        ["Current Outstanding Balance", f"{symbol} {money(remaining)}"],
        ["Current Payment Status", sale.payment_status_display],
    ]
    money_table = Table(
        [
            [Paragraph(str(a), body), Paragraph(str(b), body)]
            for a, b in money_rows
        ],
        colWidths=[60 * mm, 60 * mm],
    )
    money_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7C7BC")),
        ("BACKGROUND", (0, 3), (1, 3), colors.HexColor("#EAF4EE")),
        ("BACKGROUND", (0, 2), (1, 2), colors.HexColor("#EAF4EE")),
        ("FONTNAME", (0, 2), (1, 2), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [money_table, Spacer(1, 26)]

    sign_table = Table(
        [
            ["_________________________", "", "_________________________"],
            ["Customer signature", "", "Authorised signature / stamp"],
        ],
        colWidths=[70 * mm, 20 * mm, 70 * mm],
    )
    sign_table.setStyle(TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(sign_table)
    story.append(Spacer(1, 12))
    story.append(
        Paragraph(
            "Thank you for your payment. This receipt is computer generated.",
            styles["meta"],
        )
    )

    doc.build(story)
    buffer.seek(0)

    response = HttpResponse(buffer.read(), content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="receipt_{clean_filename(payment.receipt_number)}.pdf"'
    )
    return response